"""文件解析：PDF / DOCX / XLSX / CSV / HTML / TXT / MD / CNKI 导出。返回 (text, meta)。"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

CNKI_FIELD = re.compile(r"^([A-Za-z]+)-([^:：]+)[:：]\s*(.*)$")


def parse_pdf(path: Path) -> tuple[str, dict[str, Any]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            parts.append("")
    text = "\n\n".join(parts)
    meta: dict[str, Any] = {"pages": len(reader.pages)}
    info = reader.metadata or {}
    if info.get("/Title"):
        meta["title"] = str(info.get("/Title"))
    if info.get("/Author"):
        meta["authors"] = str(info.get("/Author"))
    return _normalize(text), meta


def parse_docx(path: Path) -> tuple[str, dict[str, Any]]:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for tb in d.tables:
        for row in tb.rows:
            parts.append(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
    title = ""
    if d.core_properties and d.core_properties.title:
        title = d.core_properties.title
    return _normalize("\n".join(parts)), {"title": title}


def parse_table(path: Path) -> tuple[str, dict[str, Any]]:
    import pandas as pd

    if path.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        sheets = pd.read_excel(path, sheet_name=None)
    else:
        sheets = {"csv": pd.read_csv(path)}
    parts = []
    for name, df in sheets.items():
        parts.append(f"## 工作表 {name}")
        parts.append(df.to_csv(index=False, sep="|"))
    return _normalize("\n".join(parts)), {"sheets": list(sheets)}


def parse_html(path: Path) -> tuple[str, dict[str, Any]]:
    html = path.read_text(encoding="utf-8", errors="ignore")
    try:
        import trafilatura

        text = trafilatura.extract(html, include_tables=True) or ""
        md = trafilatura.extract_metadata(html)
        title = (md.title if md else "") or ""
    except Exception:  # noqa: BLE001
        text, title = "", ""
    if not text:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(html, "lxml")
        text = soup.get_text("\n")
        title = soup.title.string if soup.title and soup.title.string else ""
    return _normalize(text), {"title": title}


def parse_text(path: Path) -> tuple[str, dict[str, Any]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            return _normalize(raw.decode(enc)), {}
        except UnicodeDecodeError:
            continue
    return _normalize(raw.decode("utf-8", errors="ignore")), {}


def parse_cnki_export(text: str) -> list[dict[str, str]]:
    """解析 CNKI / 万方“自定义”或 RefWorks 风格导出（形如 `Title-题名: xxx`）。返回记录列表。"""
    records: list[dict[str, str]] = []
    cur: dict[str, str] = {}
    last_key = ""
    for line in text.splitlines():
        if not line.strip():
            if cur:
                records.append(cur)
                cur, last_key = {}, ""
            continue
        m = CNKI_FIELD.match(line.strip())
        if m:
            key = m.group(1).lower()
            cur[key] = m.group(3).strip()
            last_key = key
        elif last_key:
            cur[last_key] += " " + line.strip()
    if cur:
        records.append(cur)
    return [r for r in records if r.get("title")]


def is_cnki_export(text: str) -> bool:
    head = "\n".join(text.splitlines()[:30])
    return bool(re.search(r"^(Title|SrcDatabase|Author)-", head, flags=re.M))


def parse_file(path: Path) -> tuple[str, dict[str, Any]]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext in (".xlsx", ".xlsm", ".xls", ".csv"):
        return parse_table(path)
    if ext in (".html", ".htm"):
        return parse_html(path)
    return parse_text(path)


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t　]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    # 修复 PDF 常见的英文断词连字符
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    return text.strip()
