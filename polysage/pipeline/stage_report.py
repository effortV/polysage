"""⑤ 报告与首轮实验意见：汇总 ①～④ 产出 → 首轮试验建议（LLM 选代表 + DOE 补点）→ 内部版 / 甲方版 Word。"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Callable

import pandas as pd
from docx import Document
from docx.shared import Pt

from .. import db, kb, llm
from ..config import KB
from ..formulation import library as L
from ..formulation.materials import list_materials
from . import stage_experiment, state, task

PICK_PROMPT = (
    "任务：{brief}\n\n下面是 Top 20 候选配方（含成本、四项定性预期、风险、过关把握）。首轮小试只能安排 6～8 个配方，"
    "要求覆盖不同设计思路（如茂金属置换 / 再生线性料 / EVA 保热封 / 去 HDPE / 温和版），兼顾“降本大”和“把握高”。"
    "输出 JSON：{{\"picks\": [{{\"rank\": 排名, \"reason\": \"入选理由\", \"watch\": \"现场与检测要重点看的项\"}}], "
    "\"advice\": \"给甲方的首轮试验总体建议（工艺固定、留样、再生料同批次等，3～5 句）\"}}\n\nTop 20：\n{top}"
)


def _md_to_docx(doc: Document, md: str) -> None:
    table_buf: list[list[str]] = []

    def flush_table():
        nonlocal table_buf
        rows = [r for r in table_buf if not re.fullmatch(r"[\s|:\-]+", "".join(r))]
        if rows:
            tb = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            tb.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    tb.cell(i, j).text = c.strip()
                    for p in tb.cell(i, j).paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(9)
        table_buf = []

    for line in md.splitlines():
        s = line.rstrip()
        if s.startswith("|"):
            table_buf.append([c for c in s.strip("|").split("|")])
            continue
        if table_buf:
            flush_table()
        if not s.strip():
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            level = min(len(m.group(1)), 3)
            doc.add_heading(m.group(2).strip(), level=level)
        elif re.match(r"^\s*[-*]\s+", s):
            doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", s), style="List Bullet")
        elif re.match(r"^\s*\d+[.、)]\s+", s):
            doc.add_paragraph(re.sub(r"^\s*\d+[.、)]\s+", "", s), style="List Number")
        else:
            doc.add_paragraph(s.replace("**", ""))
    if table_buf:
        flush_table()


def _df_table(doc: Document, df: pd.DataFrame, font: int = 8) -> None:
    if df.empty:
        doc.add_paragraph("（无）")
        return
    tb = doc.add_table(rows=len(df) + 1, cols=len(df.columns))
    tb.style = "Table Grid"
    for j, c in enumerate(df.columns):
        tb.cell(0, j).text = str(c)
    for i, (_, r) in enumerate(df.iterrows(), 1):
        for j, c in enumerate(df.columns):
            v = r[c]
            tb.cell(i, j).text = "" if (v is None or (isinstance(v, float) and pd.isna(v))) else str(v)
    for row in tb.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(font)


def pick_first_round(top: list[dict[str, Any]], brief: str) -> dict[str, Any]:
    listing = "\n".join(f"{r['rank']}. {r['formula']} | 成本 {r['cost']} | 降本 {r['savings_pct']}% | 预期 {r.get('expected')} | 风险 {r.get('risk')} | 把握 {r.get('confidence')}"
                        for r in top)
    try:
        data = llm.chat_json([{"role": "system", "content": "你是项目技术负责人，输出严格 JSON。"},
                              {"role": "user", "content": PICK_PROMPT.format(brief=brief, top=listing)}], max_tokens=1500, thinking=1024)
    except llm.LLMNotConfigured:
        picks = [{"rank": r["rank"], "reason": "（LLM 未配置：按排名取前 6）", "watch": ""} for r in top[:6]]
        return {"picks": picks, "advice": "工艺固定、同批再生料、留样；LLM 配置后可重新生成建议。"}
    return data if isinstance(data, dict) else {"picks": [], "advice": ""}


def build_docs(echo: Callable[[str], None] | None = None) -> list[Path]:
    t = task.load()
    brief = task.brief(t)
    design = state.stage("design")
    top = design.get("top", [])
    collect = state.stage("collect")
    mech_md = kb.read_text(KB["project"] / "现配方机理报告.md") or "（机理报告尚未生成）"
    mats = list_materials(active_only=False)
    picks = pick_first_round(top, brief) if top else {"picks": [], "advice": ""}
    # 首轮 DOE（优先配方 = 入选者）
    pick_ranks = {int(p.get("rank", 0)) for p in picks.get("picks", [])}
    priority = []
    for r in top:
        if r["rank"] in pick_ranks:
            f = _formulation_by_formula(r["formula"])
            if f:
                priority.append(f["components"])
    doe_path = None
    try:
        doe_path = stage_experiment.doe_round1(priority=priority, echo=echo)
    except Exception as e:  # noqa: BLE001
        state.log(f"DOE 生成失败：{e}", echo)
    outputs = []
    for version in ("内部版", "甲方版"):
        doc = Document()
        doc.add_heading(f"包装膜配方降本 —— 发现阶段报告（{version}）", 0)
        doc.add_paragraph(f"生成时间：{db.now()}    工具：膜方 AI（DeepSeek-V4-Pro 驱动）")
        doc.add_heading("1 任务书", 1)
        doc.add_paragraph(brief)
        doc.add_heading("2 资料采集概况", 1)
        doc.add_paragraph(f"检索命中 {collect.get('n_hits', 0)} 条，去重 {collect.get('n_unique', 0)} 条，相关保留 {collect.get('n_kept', 0)} 条，"
                          f"入库 {collect.get('n_ingested', 0)} 条，生成卡片 {collect.get('n_cards', 0)} 张。来源清单见附件《来源清单.xlsx》。")
        st = db.q("SELECT source_type, provider, COUNT(*) n FROM sources GROUP BY source_type, provider ORDER BY n DESC")
        if st:
            _df_table(doc, pd.DataFrame(st).rename(columns={"source_type": "类型", "provider": "来源库", "n": "数量"}))
        doc.add_heading("3 现配方机理分析", 1)
        _md_to_docx(doc, re.sub(r"^# .*\n", "", mech_md))
        doc.add_heading("4 候选材料清单", 1)
        mdf = pd.DataFrame([{"代码": m["code"], "材料": m["name"], "作用": m.get("role"), "典型用量%": f"{m.get('typical_min')}～{m.get('typical_max')}",
                             "拉/撕/穿/封": " ".join(str((m.get('effects') or {}).get(k, '?'))[:3] for k in ("拉伸", "撕裂", "穿刺", "热封")),
                             "参考价(元/吨)": (m.get("price_estimate") if version == "内部版" else ("实价" if m.get("price_actual") else "待询价")),
                             "风险": m.get("risk"), "用于配方": m.get("use_flag")} for m in mats])
        _df_table(doc, mdf)
        if version == "内部版":
            doc.add_paragraph("参考价为网查/假设值，仅用于排序；正式成本以甲方回填实价为准。")
        doc.add_heading("5 Top 20 候选配方【初步预计】", 1)
        tdf = pd.DataFrame([{"排名": r["rank"], "配方（组分+比例%）": r["formula"],
                             "估算成本(元/吨)": (r["cost"] if version == "内部版" else "待核算"),
                             "预期影响说明": r.get("expected"), "风险点": r.get("risk"), "过关把握": r.get("confidence")} for r in top])
        _df_table(doc, tdf)
        doc.add_paragraph("四项性能为定性预期（↑↑/↑/≈/↓/↓↓），精确值以同工艺实测为准；配方合计 100%，助剂在表内。")
        doc.add_heading("6 首轮实验意见", 1)
        doc.add_paragraph(picks.get("advice", ""))
        pdf = pd.DataFrame([{"排名": p.get("rank"), "配方": next((r["formula"] for r in top if r["rank"] == p.get("rank")), ""),
                             "入选理由": p.get("reason"), "重点观察": p.get("watch")} for p in picks.get("picks", [])])
        _df_table(doc, pdf)
        if doe_path:
            doc.add_paragraph(f"用于建模的补点设计见《首轮试验方案.xlsx》（共 {stage_experiment.count_design()} 个配方，含上述优先配方）。"
                              "若首轮无配方过关，模型仍能从“不过关”数据学到边界，第二轮向 base 方向收缩。")
        doc.add_heading("7 下一步", 1)
        for s in ["甲方回填《询价清单.xlsx》到厂含税价 → 重算 Top 20（V2）",
                  "确认 40% 内部比例、现用牌号、膜厚与工艺 → 更新任务书",
                  "先测现用膜 base（四项 + 厚度），再按《首轮试验方案》同机台同工艺打样",
                  "实测数据按《数据表模板.csv》回灌 → 建模 → 扫描候选空间 → 下一轮推荐"]:
            doc.add_paragraph(s, style="List Number")
        if version == "内部版":
            doc.add_heading("8 出处清单", 1)
            srcs = db.q("SELECT id, title, provider, year, doi, url FROM sources ORDER BY id")
            _df_table(doc, pd.DataFrame(srcs).rename(columns={"id": "#", "title": "题名", "provider": "来源库", "year": "年份", "doi": "DOI", "url": "链接"}), font=7)
        out = KB["project"] / f"发现阶段报告_{version}.docx"
        doc.save(out)
        outputs.append(out)
    return outputs


def _formulation_by_formula(formula: str) -> dict[str, Any] | None:
    for f in L.list_formulations():
        if f["formula"] == formula:
            return f
    return None


def run(echo: Callable[[str], None] | None = None) -> dict[str, Any]:
    with state.running("report", echo):
        outs = build_docs(echo)
        summary = {"outputs": [str(p) for p in outs]}
        state.set_stage("report", **summary)
    return summary
