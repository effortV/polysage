"""把对话回答（Markdown）转成 Word，并从回答里认出它生成的文件，供界面直接下载。"""
from __future__ import annotations

import re
from pathlib import Path

from .config import DATA_DIR, KB, KB_DIR

EXPORT_DIR = DATA_DIR / "exports"
FILE_EXTS = (".xlsx", ".xls", ".csv", ".docx", ".md", ".json", ".pdf")
_PATH_RE = re.compile(r"[0-9A-Za-z\u4e00-\u9fff_\-./\\]+\.(?:xlsx|xls|csv|docx|md|json|pdf)")


def referenced_files(text: str, limit: int = 8) -> list[Path]:
    """从回答正文里找出它提到的产出文件（如 `05_配方库/智能体方案_2026-09-24T0321.xlsx`），只返回真实存在的。"""
    out: list[Path] = []
    seen: set[Path] = set()
    for raw in _PATH_RE.findall(text or ""):
        cand = raw.strip("`'\"（）()，,。;；").replace("\\", "/")
        for p in _candidates(cand):
            try:
                if p.is_file() and p not in seen:
                    seen.add(p)
                    out.append(p)
                    break
            except OSError:
                continue
        if len(out) >= limit:
            break
    return out


def _candidates(rel: str) -> list[Path]:
    p = Path(rel)
    cands = [p] if p.is_absolute() else []
    cands.append(KB_DIR / rel)                      # knowledge/05_配方库/xxx.xlsx
    cands.append(DATA_DIR.parent / rel)             # 项目根下的相对路径
    name = p.name
    for d in KB.values():                           # 只给了文件名时，在各知识库目录里找
        cands.append(d / name)
    return cands


# ---------------- Markdown → Word ----------------

def answer_to_docx(text: str, path: Path | None = None, title: str = "膜方 · 对话回答") -> Path:
    """把回答转成 .docx：标题、段落、无序/有序列表、Markdown 表格。"""
    from docx import Document
    from docx.shared import Pt

    path = path or EXPORT_DIR / f"回答_{_stamp()}.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=0)
    lines = (text or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if _is_table_row(line) and i + 1 < len(lines) and _is_table_sep(lines[i + 1]):
            rows, i = _collect_table(lines, i)
            _write_table(doc, rows)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            doc.add_heading(_plain(m.group(2)), level=min(len(m.group(1)), 4))
        elif re.match(r"^\s*[-*+]\s+", line):
            doc.add_paragraph(_plain(re.sub(r"^\s*[-*+]\s+", "", line)), style="List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", line):
            doc.add_paragraph(_plain(re.sub(r"^\s*\d+[.)]\s+", "", line)), style="List Number")
        elif set(line.strip()) <= {"-", "—", "=", "*"} and len(line.strip()) >= 3:
            pass                                    # 分隔线
        else:
            doc.add_paragraph(_plain(line))
        i += 1
    for p in doc.paragraphs:
        for r in p.runs:
            if not r.font.size:
                r.font.size = Pt(10.5)
    doc.save(path)
    return path


def _stamp() -> str:
    from datetime import datetime

    return datetime.now().strftime("%Y%m%d-%H%M%S")


def _plain(s: str) -> str:
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"`([^`]*)`", r"\1", s)
    return re.sub(r"\[(.+?)\]\((.+?)\)", r"\1（\2）", s).strip()


def _is_table_row(line: str) -> bool:
    return line.strip().startswith("|") and line.count("|") >= 2


def _is_table_sep(line: str) -> bool:
    return bool(re.match(r"^\s*\|?[\s:\-|]+\|[\s:\-|]*$", line)) and "-" in line


def _collect_table(lines: list[str], i: int) -> tuple[list[list[str]], int]:
    rows = [_split_row(lines[i])]
    i += 2                                          # 跳过分隔行
    while i < len(lines) and _is_table_row(lines[i]):
        rows.append(_split_row(lines[i]))
        i += 1
    return rows, i


def _split_row(line: str) -> list[str]:
    return [_plain(c) for c in line.strip().strip("|").split("|")]


def _write_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Light Grid Accent 1"
    for ri, row in enumerate(rows):
        for ci in range(ncol):
            table.cell(ri, ci).text = row[ci] if ci < len(row) else ""
    for cell in table.rows[0].cells:                # 表头加粗
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
